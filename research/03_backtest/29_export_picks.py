"""研究步骤 29：把美人肩候选导出成表格（CSV + 可选 Excel）。

输出字段包含作战参数：止损线、二次涨停目标位、盈亏比、风险标记、推荐等级。

用法::

    python research/03_backtest/29_export_picks.py
    python research/03_backtest/29_export_picks.py --top 50 --out picks.csv
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import os
import sys
from typing import Any, Dict, List

_p = os.path.dirname(os.path.abspath(__file__))
while not os.path.isdir(os.path.join(_p, "stock_agent")) and os.path.dirname(_p) != _p:
    _p = os.path.dirname(_p)
sys.path.insert(0, _p)

from stock_agent.screener.patterns import SecondBoardConfig  # noqa: E402
from stock_agent.screener.scanner import Screener  # noqa: E402

ROOT = _p

HEADERS = [
    "排名", "推荐等级", "代码", "名称", "打分", "二次涨停概率%",
    "现价", "今日涨跌%", "蓄势天数", "60日涨停", "量能趋势",
    "总振幅%", "距锚点%", "收盘位置", "首板日", "锚点价",
    "止损线", "目标位", "盈亏比", "风险提示",
]


def grade(score: float, prob: float) -> str:
    if score >= 45 and prob >= 10.5:
        return "A（重点）"
    if score >= 30 and prob >= 9.5:
        return "B（推荐）"
    if score >= 15:
        return "C（观察）"
    return "D（弱）"


def risks(rec: Dict[str, Any]) -> str:
    out: List[str] = []
    dev = float(rec.get("距首板价%") or 0)
    gap = int(rec.get("整理天数") or 0)
    price = float(rec.get("现价") or 0)
    score = float(rec.get("score") or 0)
    if dev <= -6.0:
        out.append("贴近下沿易破位")
    if dev >= 6.0:
        out.append("已冲到上沿")
    if gap >= 12:
        out.append("蓄势过长")
    if price >= 50:
        out.append("高价股难封板")
    if score < 10:
        out.append("打分偏低")
    return "；".join(out) if out else "—"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--top", type=int, default=60)
    ap.add_argument("--market", default="main")
    ap.add_argument("--out", default=os.path.join(ROOT, ".data", "美人肩候选_最新.csv"))
    ap.add_argument("--xlsx", action="store_true")
    args = ap.parse_args()

    async def run() -> List[Dict[str, Any]]:
        cfg = SecondBoardConfig()
        async with Screener(concurrency=16) as sc:
            rows = await sc.scan_second_board(cfg, top=args.top, market=args.market, adjust="qfq")
            date = (sc.store.recent_dates("qfq", 1) or [""])[0]
            print(f"数据日期 {date}，命中 {len(rows)} 只")
            return rows

    rows = asyncio.run(run())
    if not rows:
        print("没有候选")
        return 0

    band = SecondBoardConfig().close_band
    table: List[List[Any]] = []
    for i, r in enumerate(rows, 1):
        price = float(r.get("现价") or 0)
        anchor = float(r.get("首板价") or 0)
        stop = round(anchor * (1 - band), 2)
        target = round(price * 1.10, 2)
        risk = max(price - stop, 0.01)
        rr = round((target - price) / risk, 2)
        table.append([
            i, grade(float(r.get("score") or 0), float(r.get("prob") or 0) * 100),
            r.get("code"), r.get("name"), r.get("score"),
            round(float(r.get("prob") or 0) * 100, 1),
            price, r.get("今日涨跌%"), r.get("整理天数"), r.get("60日涨停次数"),
            r.get("量能趋势"), round(float(r.get("range_pct") or 0) * 100, 1),
            r.get("距首板价%"), r.get("收盘位置"), r.get("首板日"), anchor,
            stop, target, rr, risks(r),
        ])

    out = args.out
    os.makedirs(os.path.dirname(out), exist_ok=True)
    # utf-8-sig：Excel 双击直接正常显示中文
    with open(out, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(HEADERS)
        w.writerows(table)
    print(f"已写入 CSV: {out}")

    if args.xlsx:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment

            wb = Workbook()
            ws = wb.active
            ws.title = "美人肩候选"
            ws.append(HEADERS)
            for row in table:
                ws.append(row)
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="C0504D")
                c.alignment = Alignment(horizontal="center")
            widths = [6, 12, 9, 11, 7, 15, 9, 11, 9, 10, 9, 10, 10, 10, 12, 9, 9, 9, 8, 22]
            for i, wd in enumerate(widths, 1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = wd
            ws.freeze_panes = "A2"
            xp = out.replace(".csv", ".xlsx")
            wb.save(xp)
            print(f"已写入 Excel: {xp}")
        except ImportError:
            print("（未安装 openpyxl，跳过 xlsx；pip install openpyxl 可启用）")

    # 控制台预览 A/B 级
    print("\n推荐等级 A / B 的标的：")
    print("-" * 96)
    for row in table:
        if row[1].startswith(("A", "B")):
            print(f"{row[0]:>2}. {row[2]} {row[3]:<6} 概率{row[5]:>5}% 分{row[4]:>5} "
                  f"现价{row[6]:>7} 蓄势{row[8]}天 止损{row[16]:>7} 目标{row[17]:>7} "
                  f"盈亏比{row[18]:>4} {row[19]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
